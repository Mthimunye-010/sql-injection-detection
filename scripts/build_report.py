"""Build the NCOS734 PG4 mini-project report as a .docx.

Matches the module template's page setup (A4, template margins) and typography
(Arial 10 pt body), and carries a departmental footer with page numbers.

Every figure comes from results/figures/ and every number in the prose comes
from results/tables/.

    python scripts/build_report.py
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGDIR = os.path.join(ROOT, "results", "figures")
# Allow an override so the report can be rebuilt while the previous copy is
# locked by a preview handler or an open Word window.
OUT = os.environ.get("REPORT_OUT") or os.path.join(ROOT, "NCOS734_PG4_Report.docx")

FONT = "Arial"
BODY_PT = 10
SMALL_PT = 9
CODE_PT = 8
CONTENT_EMU = Emu(int((11906 - 1152 - 1152) / 1440 * 914400))   # usable width
EN, EM = "–", "—"


# --------------------------------------------------------------------------- #
# primitives
# --------------------------------------------------------------------------- #

def _style_run(run, size=BODY_PT, bold=False, italic=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    # east-asian font must be set too or Word substitutes
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), font)


def para(doc, text="", *, size=BODY_PT, bold=False, italic=False,
         align="justify", after=6, before=0, font=FONT, line=1.08,
         indent_left=None, hanging=None):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if indent_left is not None:
        pf.left_indent = Pt(indent_left)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    if text:
        _style_run(p.add_run(text), size, bold, italic, font)
    return p


def rich(doc, frags, **kw):
    """frags: list of (text, bold, italic)."""
    kw.setdefault("after", 6)
    p = para(doc, "", **kw)
    for t, b, i in frags:
        _style_run(p.add_run(t), kw.get("size", BODY_PT), b, i)
    return p


def heading(doc, text, marks=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    _style_run(p.add_run(text), 12, bold=True)
    if marks:
        pf.tab_stops.add_tab_stop(CONTENT_EMU, WD_TAB_ALIGNMENT.RIGHT)
        _style_run(p.add_run("\t[%s]" % marks), 12, bold=True)
    return p


def sub(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    _style_run(p.add_run(text), BODY_PT, bold=True)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    _style_run(p.add_run(text), CODE_PT, font="Consolas")
    return p


def eq(doc, text):
    p = para(doc, text, size=BODY_PT, italic=True, align="center",
             after=5, before=5, font="Cambria Math")
    return p


def caption(doc, text):
    return para(doc, text, size=SMALL_PT, italic=True, align="center",
                after=10, before=3)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    _style_run(p.add_run(text), BODY_PT)
    return p


def pagebreak(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- #
# tables and figures
# --------------------------------------------------------------------------- #

def _shade(cell, hexfill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcpr.append(shd)


def _borders(tbl):
    """Horizontal rules only: a top and bottom rule plus light row separators.

    The children of w:tblBorders are a fixed sequence
    (top, left, bottom, right, insideH, insideV), and w:tblBorders itself has a
    fixed position within w:tblPr -- getting either wrong yields a document Word
    offers to "repair".
    """
    spec = [("top", "single", 8, "808080"),
            ("left", "none", 0, "auto"),
            ("bottom", "single", 8, "808080"),
            ("right", "none", 0, "auto"),
            ("insideH", "single", 4, "C8C8C8"),
            ("insideV", "none", 0, "auto")]
    borders = OxmlElement("w:tblBorders")
    for edge, val, sz, col in spec:
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), val)
        if val != "none":
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), col)
        borders.append(el)
    tbl._tbl.tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription")


def table(doc, headers, rows, weights, size=SMALL_PT):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _borders(tbl)

    total = sum(weights)
    widths = [Emu(int(CONTENT_EMU * w / total)) for w in weights]

    def fill(cells, values, bold=False, shade=None):
        for i, (c, v) in enumerate(zip(cells, values)):
            c.width = widths[i]
            p = c.paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                           else WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            _style_run(p.add_run(str(v)), size, bold=bold)
            if shade:
                _shade(c, shade)

    fill(tbl.rows[0].cells, headers, bold=True, shade="EDEDED")
    for r in rows:
        fill(tbl.add_row().cells, r)
    return tbl


def figure(doc, filename, width_in, cap):
    path = os.path.join(FIGDIR, filename)
    if not os.path.exists(path):
        print("  WARNING missing figure: %s" % filename)
        para(doc, "[figure %s not found]" % filename, italic=True, align="center")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(path, width=Pt(width_in * 72))
    caption(doc, cap)


# --------------------------------------------------------------------------- #
# page setup and footer
# --------------------------------------------------------------------------- #

def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Emu(7560000), Emu(10692000)      # A4
    for attr, dxa in (("top_margin", 1008), ("right_margin", 1152),
                      ("bottom_margin", 900), ("left_margin", 1152),
                      ("header_distance", 706), ("footer_distance", 706)):
        setattr(s, attr, Emu(int(dxa / 1440 * 914400)))

    # python-docx's stock template emits <w:zoom/> without the required
    # w:percent attribute, which fails XSD validation; give it a value.
    for z in doc.settings.element.findall(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")

    # Normal style -> Arial 10pt
    n = doc.styles["Normal"]
    n.font.name = FONT
    n.font.size = Pt(BODY_PT)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # Footer: department | lecturer | Page N
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f.paragraph_format.tab_stops.add_tab_stop(CONTENT_EMU, WD_TAB_ALIGNMENT.RIGHT)
    _style_run(f.add_run("Computer Science & IT Department"), 8)
    _style_run(f.add_run("	Lecturer: MA Arasomwan  |  Page "), 8)
    # A PAGE field must wrap its own run, or Word renders nothing for it.
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), FONT)
    rf.set(qn("w:hAnsi"), FONT)
    szel = OxmlElement("w:sz")
    szel.set(qn("w:val"), "16")
    rpr.append(rf)
    rpr.append(szel)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(rpr)
    r.append(t)
    fld.append(r)
    f._p.append(fld)
